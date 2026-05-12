from datetime import date, datetime, timedelta
from pathlib import Path
from collections import defaultdict

from flask import Flask, redirect, render_template_string, request
from flask_sqlalchemy import SQLAlchemy
from openpyxl import load_workbook


db = SQLAlchemy()


class Customer(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), unique=True, nullable=False)
    active = db.Column(db.Boolean, default=True)


class Department(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), unique=True, nullable=False)
    active = db.Column(db.Boolean, default=True)


class ReasonCode(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), unique=True, nullable=False)
    active = db.Column(db.Boolean, default=True)


class RMA(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    rma_number = db.Column(db.String(40), unique=True, nullable=False)
    date_opened = db.Column(db.Date, default=date.today)
    order_number = db.Column(db.String(120))
    work_order_number = db.Column(db.String(120))
    customer_id = db.Column(db.Integer, db.ForeignKey('customer.id'))
    department_id = db.Column(db.Integer, db.ForeignKey('department.id'))
    reason_code_id = db.Column(db.Integer, db.ForeignKey('reason_code.id'))
    part_number = db.Column(db.String(120))
    quantity_affected = db.Column(db.Integer, default=0)
    defect_description = db.Column(db.Text)
    status = db.Column(db.String(80), default='Open')
    customer = db.relationship('Customer')
    department = db.relationship('Department')
    reason_code = db.relationship('ReasonCode')


class WorkOrder(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    work_order_number = db.Column(db.String(120), unique=True, nullable=False)
    customer_id = db.Column(db.Integer, db.ForeignKey('customer.id'))
    part_number = db.Column(db.String(120), nullable=False)
    description = db.Column(db.String(240))
    quantity_ordered = db.Column(db.Integer, default=0)
    quantity_completed = db.Column(db.Integer, default=0)
    quantity_failed = db.Column(db.Integer, default=0)
    release_date = db.Column(db.Date)
    due_date = db.Column(db.Date)
    status = db.Column(db.String(80), default='Open')
    priority = db.Column(db.String(40), default='Normal')
    material_status = db.Column(db.String(80), default='Not Reviewed')
    inventory_strategy = db.Column(db.String(80), default='Unset')
    customer = db.relationship('Customer')


class OperationClockSummary(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    work_order_id = db.Column(db.Integer, db.ForeignKey('work_order.id'))
    department_id = db.Column(db.Integer, db.ForeignKey('department.id'))
    operation_name = db.Column(db.String(160))
    period_date = db.Column(db.Date, default=date.today)
    planned_qty = db.Column(db.Integer, default=0)
    first_pass_good_qty = db.Column(db.Integer, default=0)
    rework_qty = db.Column(db.Integer, default=0)
    scrap_qty = db.Column(db.Integer, default=0)
    labor_hours = db.Column(db.Float, default=0)
    overtime_hours = db.Column(db.Float, default=0)
    work_order = db.relationship('WorkOrder')
    department = db.relationship('Department')


class OperatorThroughput(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    period_date = db.Column(db.Date, default=date.today)
    operator_label = db.Column(db.String(120), nullable=False)
    department_id = db.Column(db.Integer, db.ForeignKey('department.id'))
    operation_name = db.Column(db.String(160), nullable=False)
    part_number = db.Column(db.String(120))
    part_family = db.Column(db.String(160))
    complexity = db.Column(db.String(40), default='Medium')
    quantity_completed = db.Column(db.Integer, default=0)
    quantity_rejected = db.Column(db.Integer, default=0)
    labor_hours = db.Column(db.Float, default=0)
    setup_hours = db.Column(db.Float, default=0)
    notes = db.Column(db.Text)
    department = db.relationship('Department')


class PurchaseRequirement(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    item_number = db.Column(db.String(120), nullable=False)
    description = db.Column(db.String(240))
    work_order_id = db.Column(db.Integer, db.ForeignKey('work_order.id'))
    required_qty = db.Column(db.Float, default=0)
    on_hand_qty = db.Column(db.Float, default=0)
    supplier = db.Column(db.String(180))
    need_by_date = db.Column(db.Date)
    lead_time_days = db.Column(db.Integer, default=0)
    status = db.Column(db.String(80), default='Needs Review')
    work_order = db.relationship('WorkOrder')


class AttendanceMoraleMetric(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    period_date = db.Column(db.Date, default=date.today)
    department_id = db.Column(db.Integer, db.ForeignKey('department.id'))
    scheduled_headcount = db.Column(db.Integer, default=0)
    present_headcount = db.Column(db.Integer, default=0)
    overtime_hours = db.Column(db.Float, default=0)
    employees_over_50_hours = db.Column(db.Integer, default=0)
    exhausted_pto_count = db.Column(db.Integer, default=0)
    pto_absence_count = db.Column(db.Integer, default=0)
    sick_absence_count = db.Column(db.Integer, default=0)
    turnover_count = db.Column(db.Integer, default=0)
    staffing_shortage_count = db.Column(db.Integer, default=0)
    department = db.relationship('Department')


class OperatingProfile(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    company_name = db.Column(db.String(180))
    employee_count = db.Column(db.Integer, default=0)
    sku_count = db.Column(db.Integer, default=0)
    product_complexity = db.Column(db.String(40), default='Medium')
    demand_variability = db.Column(db.String(40), default='Medium')
    cash_flow_priority = db.Column(db.String(40), default='Medium')
    preferred_inventory_mode = db.Column(db.String(80), default='Unsure')
    current_mode = db.Column(db.String(80), default='Unknown')
    recommended_mode = db.Column(db.String(160))
    recommendation_notes = db.Column(db.Text)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow)


class ActivityLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    area = db.Column(db.String(80))
    action = db.Column(db.String(120))
    detail = db.Column(db.Text)


CSS = """body{margin:0;background:#0b1117;color:#e8eef5;font:14px Segoe UI,Arial}aside{position:fixed;inset:0 auto 0 0;width:220px;background:#080d12;border-right:1px solid #263647;padding:18px}.brand{font-size:22px;font-weight:700;margin-bottom:16px;color:#fff}nav a{display:block;color:#95a4b5;text-decoration:none;padding:8px 10px;border-radius:8px}nav a:hover{background:#172433;color:white}main{margin-left:220px;padding:20px;max-width:1500px}h1{font-size:22px;margin:0 0 16px}h2{font-size:16px}.cards{display:grid;grid-template-columns:repeat(5,minmax(135px,1fr));gap:12px}.card,section{background:#121b24;border:1px solid #263647;border-radius:12px;padding:14px}.card span{color:#95a4b5}.card b{display:block;font-size:24px;margin-top:6px}table{width:100%;border-collapse:collapse;background:#121b24;border:1px solid #263647}th,td{border-bottom:1px solid #263647;padding:8px;text-align:left;vertical-align:top}th{color:#95a4b5;background:#0f1720}input,select,textarea{box-sizing:border-box;background:#0b1117;color:#e8eef5;border:1px solid #263647;border-radius:8px;padding:8px;width:100%}textarea{min-height:85px}.form{display:grid;grid-template-columns:repeat(3,1fr);gap:12px}.wide{grid-column:1/-1}.btn,button{background:#3e78a8;color:white;border:0;border-radius:8px;padding:8px 12px;text-decoration:none;display:inline-block;cursor:pointer}.toolbar{display:flex;gap:8px;margin-bottom:12px;align-items:center;flex-wrap:wrap}.notice{border-left:4px solid #3e78a8;background:#101923;padding:10px;border-radius:8px;margin-bottom:14px;color:#c8d4df}.muted{color:#95a4b5}"""
BASE = """<!doctype html><html><head><title>ForgeQC</title><style>{{css}}</style></head><body><aside><div class='brand'>ForgeQC</div><nav><a href='/'>Dashboard</a><a href='/setup'>Setup Profile</a><a href='/workorders'>Work Orders</a><a href='/clocking'>FPY Clocking</a><a href='/efficiency'>Operator Efficiency</a><a href='/planning'>Planning</a><a href='/rma'>RMA</a><a href='/metrics'>Metrics</a><a href='/morale'>Morale Pulse</a><a href='/admin'>Admin</a></nav></aside><main><h1>{{title}}</h1>{{body|safe}}</main></body></html>"""


def page(title, body):
    return render_template_string(BASE, css=CSS, title=title, body=body)


def parse_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if v in (None, ''): return None
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y'):
        try: return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError: pass
    return None


def as_int(v, default=0):
    try: return default if v in (None, '') else int(float(v))
    except (TypeError, ValueError): return default


def as_float(v, default=0.0):
    try: return default if v in (None, '') else float(v)
    except (TypeError, ValueError): return default


def get_or_create(model, name):
    if name is None or str(name).strip() == '': return None
    clean = str(name).strip()
    row = model.query.filter(db.func.lower(model.name) == clean.lower()).first()
    if not row:
        row = model(name=clean)
        db.session.add(row)
        db.session.flush()
    return row


def ref_options(model):
    return "<option value=''></option>" + ''.join(f"<option value='{r.id}'>{r.name}</option>" for r in model.query.filter_by(active=True).order_by(model.name))


def wo_options():
    return "<option value=''></option>" + ''.join(f"<option value='{w.id}'>{w.work_order_number} | {w.part_number}</option>" for w in WorkOrder.query.order_by(WorkOrder.work_order_number))


def next_number(model, field, prefix):
    year = date.today().year
    count = model.query.filter(getattr(model, field).like(f'{prefix}-{year}-%')).count() + 1
    return f'{prefix}-{year}-{count:04d}'


def log(area, action, detail=''):
    db.session.add(ActivityLog(area=area, action=action, detail=detail))


def fpy(good, rework, scrap):
    total = (good or 0) + (rework or 0) + (scrap or 0)
    return round(((good or 0) / total) * 100, 1) if total else 0


def rate(qty, hours):
    return round((qty or 0) / (hours or 0), 2) if hours else 0


def fail_rate(failed, ordered):
    return round(((failed or 0) / (ordered or 0)) * 100, 1) if ordered else 0


def morale_score(m):
    absent = max((m.scheduled_headcount or 0) - (m.present_headcount or 0), 0)
    score = min((m.overtime_hours or 0) * 0.4, 30)
    score += (m.employees_over_50_hours or 0) * 5 + (m.exhausted_pto_count or 0) * 4
    score += (m.pto_absence_count or 0) * 3 + (m.sick_absence_count or 0) * 3
    score += (m.turnover_count or 0) * 8 + (m.staffing_shortage_count or 0) * 5 + absent * 4
    return round(score, 1)


def risk_label(score):
    return 'Critical' if score >= 60 else 'High' if score >= 35 else 'Medium' if score >= 18 else 'Low'


def order_by_date(req):
    return req.need_by_date - timedelta(days=req.lead_time_days or 0) if req.need_by_date else None


def recommend_mode(employees, skus, complexity, variability, cash, preference):
    if preference == 'Hold Inventory': return 'Inventory-buffered MRP', 'Use min/max, reorder points, cycle counts, and strategic stock buffers.'
    if cash == 'High' and employees <= 75: return 'Hybrid JIT with protected critical inventory', 'Cash matters, but small teams get crushed by zero-stock emergencies.'
    if complexity == 'High' or variability == 'High' or skus >= 250: return 'Hybrid Kanban/MRP', 'Complexity and variable demand need planning buckets, lead-time alerts, and selective buffers.'
    if cash in ('Medium', 'High'): return 'Lean JIT', 'Good fit for controlled product mix, stable routing, and cash-sensitive operations.'
    return 'Inventory-buffered MRP', 'Cash pressure is not dominant, so strategic inventory can protect delivery and quality.'


def create_app():
    app = Flask(__name__)
    root = Path(__file__).resolve().parent
    data_dir = root / 'data'
    imports_dir = data_dir / 'imports'
    data_dir.mkdir(exist_ok=True)
    imports_dir.mkdir(parents=True, exist_ok=True)
    app.config['SQLALCHEMY_DATABASE_URI'] = f"sqlite:///{data_dir / 'forgeqc.db'}"
    app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
    db.init_app(app)
    with app.app_context():
        db.create_all()
        for model, values in [(Department, ['Laser','Tube Laser','Machining','Welding','Forming','Powder','Quality','Shipping']), (ReasonCode, ['Dimensional','Wrong revision','Wrong material','Missing feature','Weld defect','Powder defect','Damage','Documentation']), (Customer, ['Internal'])]:
            for value in values: get_or_create(model, value)
        db.session.commit()

    @app.route('/')
    def dashboard():
        today = date.today()
        past = WorkOrder.query.filter(WorkOrder.status != 'Closed', WorkOrder.due_date < today).count()
        soon = WorkOrder.query.filter(WorkOrder.status != 'Closed', WorkOrder.due_date >= today, WorkOrder.due_date <= today + timedelta(days=14)).count()
        purch = PurchaseRequirement.query.filter(PurchaseRequirement.status.in_(['Needs Review','Needs Purchased'])).count()
        fpys = [fpy(c.first_pass_good_qty, c.rework_qty, c.scrap_qty) for c in OperationClockSummary.query.all()]
        avg_fpy = round(sum(fpys) / len(fpys), 1) if fpys else 0
        pulses = [morale_score(m) for m in AttendanceMoraleMetric.query.order_by(AttendanceMoraleMetric.period_date.desc()).limit(20)]
        pulse = round(sum(pulses) / len(pulses), 1) if pulses else 0
        cards = {'Past Due WO': past, 'Due Soon': soon, 'Purchasing Needs': purch, 'Avg FPY': f'{avg_fpy}%', 'Morale Pulse': f'{risk_label(pulse)} {pulse}'}
        card_html = ''.join(f"<div class='card'><span>{k}</span><b>{v}</b></div>" for k, v in cards.items())
        watch = ''.join(f"<tr><td>{w.work_order_number}</td><td>{w.part_number}</td><td>{w.due_date}</td><td>{w.material_status}</td><td>{w.status}</td></tr>" for w in WorkOrder.query.filter(WorkOrder.status != 'Closed').order_by(WorkOrder.due_date).limit(15))
        return page('Dashboard', f"<div class='cards'>{card_html}</div><br><section><h2>Work Order Watchlist</h2><table><tr><th>WO</th><th>Part</th><th>Due</th><th>Material</th><th>Status</th></tr>{watch}</table></section>")

    @app.route('/setup', methods=['GET','POST'])
    def setup():
        p = OperatingProfile.query.order_by(OperatingProfile.id.desc()).first() or OperatingProfile()
        if request.method == 'POST':
            p.company_name = request.form.get('company_name'); p.employee_count = as_int(request.form.get('employee_count')); p.sku_count = as_int(request.form.get('sku_count'))
            p.product_complexity = request.form.get('product_complexity'); p.demand_variability = request.form.get('demand_variability'); p.cash_flow_priority = request.form.get('cash_flow_priority')
            p.preferred_inventory_mode = request.form.get('preferred_inventory_mode'); p.current_mode = request.form.get('current_mode')
            p.recommended_mode, p.recommendation_notes = recommend_mode(p.employee_count, p.sku_count, p.product_complexity, p.demand_variability, p.cash_flow_priority, p.preferred_inventory_mode)
            p.updated_at = datetime.utcnow(); db.session.add(p); log('Setup','Profile saved',p.recommended_mode); db.session.commit(); return redirect('/setup')
        form = f"""<section><div class='notice'>Define how the company chooses to operate: JIT, inventory-buffered, or hybrid.</div><form method='post' class='form'><label>Company<input name='company_name' value='{p.company_name or ''}'></label><label>Employees<input name='employee_count' type='number' value='{p.employee_count or 0}'></label><label>Active Parts/SKUs<input name='sku_count' type='number' value='{p.sku_count or 0}'></label><label>Product Complexity<select name='product_complexity'><option>Low</option><option>Medium</option><option>High</option></select></label><label>Demand Variability<select name='demand_variability'><option>Low</option><option>Medium</option><option>High</option></select></label><label>Cash Flow Priority<select name='cash_flow_priority'><option>Low</option><option>Medium</option><option>High</option></select></label><label>Preferred Inventory Mode<select name='preferred_inventory_mode'><option>Unsure</option><option>JIT</option><option>Hold Inventory</option><option>Hybrid</option></select></label><label>Current Mode<input name='current_mode' value='{p.current_mode or ''}'></label><div class='wide'><button>Save And Recommend</button></div></form></section><br><section><h2>Recommendation</h2><p><b>{p.recommended_mode or 'Not calculated yet'}</b></p><p>{p.recommendation_notes or ''}</p></section>"""
        return page('Operating Setup Profile', form)

    @app.route('/workorders', methods=['GET','POST'])
    def workorders():
        if request.method == 'POST':
            c = get_or_create(Customer, request.form.get('customer'))
            w = WorkOrder(work_order_number=request.form.get('work_order_number'), customer_id=c.id if c else None, part_number=request.form.get('part_number'), description=request.form.get('description'), quantity_ordered=as_int(request.form.get('quantity_ordered')), quantity_completed=as_int(request.form.get('quantity_completed')), quantity_failed=as_int(request.form.get('quantity_failed')), release_date=parse_date(request.form.get('release_date')), due_date=parse_date(request.form.get('due_date')), status=request.form.get('status') or 'Open', priority=request.form.get('priority') or 'Normal', material_status=request.form.get('material_status') or 'Not Reviewed', inventory_strategy=request.form.get('inventory_strategy') or 'Unset')
            db.session.add(w); log('WorkOrder','Save',w.work_order_number); db.session.commit(); return redirect('/workorders')
        form = """<section><form method='post' class='form'><label>Work Order<input name='work_order_number' required></label><label>Customer<input name='customer'></label><label>Part Number<input name='part_number' required></label><label>Description<input name='description'></label><label>Qty Ordered<input name='quantity_ordered' type='number'></label><label>Qty Completed<input name='quantity_completed' type='number'></label><label>Qty Failed<input name='quantity_failed' type='number'></label><label>Release Date<input name='release_date' type='date'></label><label>Due Date<input name='due_date' type='date'></label><label>Status><select name='status'><option>Open</option><option>Running</option><option>Hold</option><option>Closed</option></select></label><label>Priority<select name='priority'><option>Normal</option><option>Hot</option><option>Critical</option></select></label><label>Material Status<select name='material_status'><option>Not Reviewed</option><option>Needs Purchased</option><option>Ordered</option><option>Available</option><option>Short</option></select></label><label>Inventory Strategy<select name='inventory_strategy'><option>Unset</option><option>JIT</option><option>Buffer Stock</option><option>Hybrid</option></select></label><button>Save Work Order</button></form></section><br>"""
        rows = ''.join(f"<tr><td>{w.work_order_number}</td><td>{w.customer.name if w.customer else ''}</td><td>{w.part_number}</td><td>{w.quantity_ordered}</td><td>{w.quantity_failed}</td><td>{fail_rate(w.quantity_failed,w.quantity_ordered)}%</td><td>{w.due_date}</td><td>{w.material_status}</td><td>{w.status}</td></tr>" for w in WorkOrder.query.order_by(WorkOrder.due_date).all())
        return page('Work Orders', form + f"<table><tr><th>WO</th><th>Customer</th><th>Part</th><th>Qty</th><th>Failed</th><th>Fail Rate</th><th>Due</th><th>Material</th><th>Status</th></tr>{rows}</table>")

    @app.route('/clocking', methods=['GET','POST'])
    def clocking():
        if request.method == 'POST':
            r = OperationClockSummary(work_order_id=as_int(request.form.get('work_order_id'), None), department_id=as_int(request.form.get('department_id'), None), operation_name=request.form.get('operation_name'), period_date=parse_date(request.form.get('period_date')) or date.today(), planned_qty=as_int(request.form.get('planned_qty')), first_pass_good_qty=as_int(request.form.get('first_pass_good_qty')), rework_qty=as_int(request.form.get('rework_qty')), scrap_qty=as_int(request.form.get('scrap_qty')), labor_hours=as_float(request.form.get('labor_hours')), overtime_hours=as_float(request.form.get('overtime_hours')))
            db.session.add(r); log('FPY','Aggregate clocking saved',r.operation_name or ''); db.session.commit(); return redirect('/clocking')
        form = f"<div class='notice'>Aggregate operation clocking only. No employee names, numbers, badge IDs, or payroll data.</div><section><form method='post' class='form'><label>Work Order<select name='work_order_id'>{wo_options()}</select></label><label>Department<select name='department_id'>{ref_options(Department)}</select></label><label>Operation<input name='operation_name'></label><label>Date<input name='period_date' type='date'></label><label>Planned Qty<input name='planned_qty' type='number'></label><label>First Pass Good<input name='first_pass_good_qty' type='number'></label><label>Rework Qty<input name='rework_qty' type='number'></label><label>Scrap Qty<input name='scrap_qty' type='number'></label><label>Labor Hours<input name='labor_hours' type='number' step='0.1'></label><label>OT Hours<input name='overtime_hours' type='number' step='0.1'></label><button>Save FPY Summary</button></form></section><br>"
        rows = ''.join(f"<tr><td>{c.period_date}</td><td>{c.work_order.work_order_number if c.work_order else ''}</td><td>{c.department.name if c.department else ''}</td><td>{c.operation_name}</td><td>{c.first_pass_good_qty}</td><td>{c.rework_qty}</td><td>{c.scrap_qty}</td><td>{fpy(c.first_pass_good_qty,c.rework_qty,c.scrap_qty)}%</td><td>{c.overtime_hours}</td></tr>" for c in OperationClockSummary.query.order_by(OperationClockSummary.period_date.desc()).limit(200))
        return page('FPY Through Clocking', form + f"<table><tr><th>Date</th><th>WO</th><th>Dept</th><th>Operation</th><th>First Pass</th><th>Rework</th><th>Scrap</th><th>FPY</th><th>OT</th></tr>{rows}</table>")

    @app.route('/efficiency', methods=['GET','POST'])
    def efficiency():
        if request.method == 'POST':
            e = OperatorThroughput(period_date=parse_date(request.form.get('period_date')) or date.today(), operator_label=request.form.get('operator_label'), department_id=as_int(request.form.get('department_id'), None), operation_name=request.form.get('operation_name'), part_number=request.form.get('part_number'), part_family=request.form.get('part_family'), complexity=request.form.get('complexity') or 'Medium', quantity_completed=as_int(request.form.get('quantity_completed')), quantity_rejected=as_int(request.form.get('quantity_rejected')), labor_hours=as_float(request.form.get('labor_hours')), setup_hours=as_float(request.form.get('setup_hours')), notes=request.form.get('notes'))
            db.session.add(e); log('Efficiency','Throughput saved',f'{e.operator_label} {e.operation_name}'); db.session.commit(); return redirect('/efficiency')
        form = f"<div class='notice'>Local operator throughput for quoting. This is separate from privacy-safe morale data.</div><section><form method='post' class='form'><label>Date<input name='period_date' type='date'></label><label>Operator Label<input name='operator_label' required></label><label>Department<select name='department_id'>{ref_options(Department)}</select></label><label>Operation<input name='operation_name' required></label><label>Part Number<input name='part_number'></label><label>Part Family<input name='part_family'></label><label>Complexity<select name='complexity'><option>Low</option><option>Medium</option><option>High</option></select></label><label>Qty Completed<input name='quantity_completed' type='number'></label><label>Qty Rejected<input name='quantity_rejected' type='number'></label><label>Labor Hours<input name='labor_hours' type='number' step='0.1'></label><label>Setup Hours<input name='setup_hours' type='number' step='0.1'></label><label class='wide'>Notes<textarea name='notes'></textarea></label><button>Save Throughput</button></form></section><br>"
        entries = ''.join(f"<tr><td>{e.operator_label}</td><td>{e.department.name if e.department else ''}</td><td>{e.operation_name}</td><td>{e.part_family or e.part_number or ''}</td><td>{e.complexity}</td><td>{e.quantity_completed}</td><td>{e.quantity_rejected}</td><td>{e.labor_hours}</td><td>{rate(e.quantity_completed,e.labor_hours)}</td></tr>" for e in OperatorThroughput.query.order_by(OperatorThroughput.period_date.desc()).limit(300))
        groups = defaultdict(list)
        for e in OperatorThroughput.query.all(): groups[(e.operation_name, e.part_family or e.part_number or 'General', e.complexity)].append(rate(e.quantity_completed, e.labor_hours))
        summary = ''.join(f"<tr><td>{k[0]}</td><td>{k[1]}</td><td>{k[2]}</td><td>{round(min(v),2)}</td><td>{round(sum(v)/len(v),2)}</td><td>{round(max(v),2)}</td></tr>" for k, v in groups.items() if v)
        return page('Operator Efficiency For Quoting', form + f"<section><h2>Throughput Entries</h2><table><tr><th>Operator</th><th>Dept</th><th>Operation</th><th>Part/Family</th><th>Complexity</th><th>Qty</th><th>Rejected</th><th>Hours</th><th>Qty/Hr</th></tr>{entries}</table></section><br><section><h2>Quoting Baselines From Real Production</h2><table><tr><th>Operation</th><th>Part/Family</th><th>Complexity</th><th>Low</th><th>Average</th><th>Best</th></tr>{summary}</table></section>")

    @app.route('/planning', methods=['GET','POST'])
    def planning():
        if request.method == 'POST':
            p = PurchaseRequirement(item_number=request.form.get('item_number'), description=request.form.get('description'), work_order_id=as_int(request.form.get('work_order_id'), None), required_qty=as_float(request.form.get('required_qty')), on_hand_qty=as_float(request.form.get('on_hand_qty')), supplier=request.form.get('supplier'), need_by_date=parse_date(request.form.get('need_by_date')), lead_time_days=as_int(request.form.get('lead_time_days')), status=request.form.get('status') or 'Needs Review')
            db.session.add(p); log('Planning','Purchase need saved',p.item_number); db.session.commit(); return redirect('/planning')
        today = date.today()
        past = ''.join(f"<tr><td>{w.work_order_number}</td><td>{w.part_number}</td><td>{w.due_date}</td><td>{w.material_status}</td></tr>" for w in WorkOrder.query.filter(WorkOrder.status != 'Closed', WorkOrder.due_date < today).order_by(WorkOrder.due_date))
        upcoming = ''.join(f"<tr><td>{w.work_order_number}</td><td>{w.part_number}</td><td>{w.due_date}</td><td>{w.material_status}</td></tr>" for w in WorkOrder.query.filter(WorkOrder.status != 'Closed', WorkOrder.due_date >= today, WorkOrder.due_date <= today + timedelta(days=30)).order_by(WorkOrder.due_date))
        purchases = ''.join(f"<tr><td>{p.item_number}</td><td>{p.required_qty}</td><td>{p.on_hand_qty}</td><td>{p.need_by_date}</td><td>{order_by_date(p)}</td><td>{p.lead_time_days}</td><td>{p.status}</td></tr>" for p in PurchaseRequirement.query.order_by(PurchaseRequirement.need_by_date).all())
        form = f"<section><h2>Purchase Requirement</h2><form method='post' class='form'><label>Item<input name='item_number' required></label><label>Description<input name='description'></label><label>Work Order<select name='work_order_id'>{wo_options()}</select></label><label>Required Qty<input name='required_qty' type='number' step='0.01'></label><label>On Hand Qty<input name='on_hand_qty' type='number' step='0.01'></label><label>Supplier<input name='supplier'></label><label>Need By<input name='need_by_date' type='date'></label><label>Lead Time Days<input name='lead_time_days' type='number'></label><label>Status<select name='status'><option>Needs Review</option><option>Needs Purchased</option><option>Ordered</option><option>Received</option></select></label><button>Save Requirement</button></form></section><br>"
        return page('Planning, Lead Time, and JIT', form + f"<section><h2>Past Due</h2><table><tr><th>WO</th><th>Part</th><th>Due</th><th>Material</th></tr>{past}</table></section><br><section><h2>Coming Up 30 Days</h2><table><tr><th>WO</th><th>Part</th><th>Due</th><th>Material</th></tr>{upcoming}</table></section><br><section><h2>Purchasing</h2><table><tr><th>Item</th><th>Required</th><th>On Hand</th><th>Need By</th><th>Order By</th><th>Lead Time</th><th>Status</th></tr>{purchases}</table></section>")

    @app.route('/rma', methods=['GET','POST'])
    def rma():
        if request.method == 'POST':
            c = get_or_create(Customer, request.form.get('customer')); d = get_or_create(Department, request.form.get('department')); rc = get_or_create(ReasonCode, request.form.get('reason'))
            r = RMA(rma_number=next_number(RMA,'rma_number','RMA'), date_opened=parse_date(request.form.get('date_opened')) or date.today(), order_number=request.form.get('order_number'), work_order_number=request.form.get('work_order_number'), customer_id=c.id if c else None, department_id=d.id if d else None, reason_code_id=rc.id if rc else None, part_number=request.form.get('part_number'), quantity_affected=as_int(request.form.get('quantity_affected')), defect_description=request.form.get('defect_description'), status=request.form.get('status') or 'Open')
            db.session.add(r); log('RMA','Save',r.rma_number); db.session.commit(); return redirect('/rma')
        form = """<section><form method='post' class='form'><label>Date<input type='date' name='date_opened'></label><label>Order<input name='order_number'></label><label>WO<input name='work_order_number'></label><label>Customer<input name='customer'></label><label>Department<input name='department'></label><label>Reason<input name='reason'></label><label>Part<input name='part_number'></label><label>Qty Affected<input name='quantity_affected' type='number'></label><label>Status<select name='status'><option>Open</option><option>Closed</option></select></label><label class='wide'>Defect<textarea name='defect_description'></textarea></label><button>Save RMA</button></form></section><br>"""
        rows = ''.join(f"<tr><td>{r.rma_number}</td><td>{r.date_opened}</td><td>{r.customer.name if r.customer else ''}</td><td>{r.part_number}</td><td>{r.quantity_affected}</td><td>{r.department.name if r.department else ''}</td><td>{r.reason_code.name if r.reason_code else ''}</td><td>{r.status}</td></tr>" for r in RMA.query.order_by(RMA.id.desc()).limit(500))
        return page('RMA', form + f"<form method='post' action='/rma/import'><button>Import Local RMA_Tracker.xlsx</button></form><br><table><tr><th>RMA</th><th>Date</th><th>Customer</th><th>Part</th><th>Qty</th><th>Dept</th><th>Reason</th><th>Status</th></tr>{rows}</table>")

    @app.route('/rma/import', methods=['POST'])
    def import_rma():
        path = imports_dir / 'RMA_Tracker.xlsx'
        if not path.exists(): log('RMA','Import failed','Local workbook not found'); db.session.commit(); return redirect('/rma')
        wb = load_workbook(path, data_only=True); imported = 0
        if 'RMA Log' in wb.sheetnames:
            for vals in wb['RMA Log'].iter_rows(min_row=2, values_only=True):
                opened, order, customer, department, _person, desc, reason = (list(vals) + [None] * 7)[:7]
                if not any([opened, order, customer, department, desc, reason]): continue
                c = get_or_create(Customer, customer); d = get_or_create(Department, department); rc = get_or_create(ReasonCode, reason)
                db.session.add(RMA(rma_number=next_number(RMA,'rma_number','RMA'), date_opened=parse_date(opened), order_number=str(order) if order else None, customer_id=c.id if c else None, department_id=d.id if d else None, reason_code_id=rc.id if rc else None, defect_description=str(desc) if desc else None, status='Open'))
                imported += 1
        log('RMA','Import',f'Imported {imported} rows from local workbook'); db.session.commit(); return redirect('/rma')

    @app.route('/morale', methods=['GET','POST'])
    def morale():
        if request.method == 'POST':
            m = AttendanceMoraleMetric(period_date=parse_date(request.form.get('period_date')) or date.today(), department_id=as_int(request.form.get('department_id'), None), scheduled_headcount=as_int(request.form.get('scheduled_headcount')), present_headcount=as_int(request.form.get('present_headcount')), overtime_hours=as_float(request.form.get('overtime_hours')), employees_over_50_hours=as_int(request.form.get('employees_over_50_hours')), exhausted_pto_count=as_int(request.form.get('exhausted_pto_count')), pto_absence_count=as_int(request.form.get('pto_absence_count')), sick_absence_count=as_int(request.form.get('sick_absence_count')), turnover_count=as_int(request.form.get('turnover_count')), staffing_shortage_count=as_int(request.form.get('staffing_shortage_count')))
            db.session.add(m); log('Morale','Aggregate input saved','Privacy-safe'); db.session.commit(); return redirect('/morale')
        form = f"<div class='notice'>Aggregate department metrics only. No employee identifiers.</div><section><form method='post' class='form'><label>Date<input name='period_date' type='date'></label><label>Department<select name='department_id'>{ref_options(Department)}</select></label><label>Scheduled Headcount<input name='scheduled_headcount' type='number'></label><label>Present Headcount<input name='present_headcount' type='number'></label><label>Overtime Hours<input name='overtime_hours' type='number' step='0.1'></label><label>Employees Over 50 Hours<input name='employees_over_50_hours' type='number'></label><label>Exhausted PTO Count<input name='exhausted_pto_count' type='number'></label><label>PTO Absence Count<input name='pto_absence_count' type='number'></label><label>Sick Absence Count<input name='sick_absence_count' type='number'></label><label>Turnover Count<input name='turnover_count' type='number'></label><label>Staffing Shortage Count<input name='staffing_shortage_count' type='number'></label><button>Save Aggregate Pulse</button></form></section><br>"
        rows = ''.join(f"<tr><td>{m.period_date}</td><td>{m.department.name if m.department else ''}</td><td>{m.overtime_hours}</td><td>{m.sick_absence_count}</td><td>{morale_score(m)}</td><td>{risk_label(morale_score(m))}</td></tr>" for m in AttendanceMoraleMetric.query.order_by(AttendanceMoraleMetric.period_date.desc()).limit(200))
        return page('Morale Pulse', form + f"<table><tr><th>Date</th><th>Dept</th><th>OT</th><th>Sick</th><th>Score</th><th>Risk</th></tr>{rows}</table>")

    @app.route('/metrics')
    def metrics():
        dept_rows = ''
        for d in Department.query.order_by(Department.name):
            clocks = OperationClockSummary.query.filter_by(department_id=d.id).all(); fpys = [fpy(c.first_pass_good_qty,c.rework_qty,c.scrap_qty) for c in clocks]
            dept_rows += f"<tr><td>{d.name}</td><td>{round(sum(fpys)/len(fpys),1) if fpys else 0}%</td><td>{RMA.query.filter_by(department_id=d.id).count()}</td></tr>"
        wo_rows = ''.join(f"<tr><td>{w.part_number}</td><td>{w.quantity_ordered}</td><td>{w.quantity_failed}</td><td>{fail_rate(w.quantity_failed,w.quantity_ordered)}%</td></tr>" for w in WorkOrder.query.order_by(WorkOrder.quantity_failed.desc()).limit(100))
        return page('Metrics', f"<section><h2>FPY and RMA by Department</h2><table><tr><th>Dept</th><th>Avg FPY</th><th>RMA Count</th></tr>{dept_rows}</table></section><br><section><h2>Work Order Failure Trends By Quantity</h2><table><tr><th>Part</th><th>Qty Ordered</th><th>Qty Failed</th><th>Fail Rate</th></tr>{wo_rows}</table></section>")

    @app.route('/admin', methods=['GET','POST'])
    def admin():
        models = {'customers': Customer, 'departments': Department, 'reason_codes': ReasonCode}
        if request.method == 'POST': get_or_create(models[request.form['table']], request.form.get('name')); db.session.commit(); return redirect('/admin')
        body = ''
        for name, model in models.items():
            rows = ''.join(f"<tr><td>{x.name}</td><td>{x.active}</td></tr>" for x in model.query.order_by(model.name))
            body += f"<section><h2>{name}</h2><form method='post' class='toolbar'><input type='hidden' name='table' value='{name}'><input name='name' placeholder='Name'><button>Add</button></form><table><tr><th>Name</th><th>Active</th></tr>{rows}</table></section><br>"
        return page('Admin', body)

    return app
