from pathlib import Path
from datetime import datetime, date
from collections import defaultdict

from flask import Flask, request, redirect, send_file, render_template_string, abort
from flask_sqlalchemy import SQLAlchemy
from openpyxl import load_workbook, Workbook


db = SQLAlchemy()


class Customer(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), unique=True, nullable=False)
    description = db.Column(db.Text)
    active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)


class Department(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), unique=True, nullable=False)
    description = db.Column(db.Text)
    active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)


class ReasonCode(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), unique=True, nullable=False)
    description = db.Column(db.Text)
    active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)


class RMA(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    rma_number = db.Column(db.String(32), unique=True, nullable=False)
    date_opened = db.Column(db.Date, default=date.today)
    customer_id = db.Column(db.Integer, db.ForeignKey("customer.id"))
    order_number = db.Column(db.String(120))
    part_number = db.Column(db.String(120))
    quantity_affected = db.Column(db.Integer)
    department_id = db.Column(db.Integer, db.ForeignKey("department.id"))
    responsible_area = db.Column(db.String(180))
    defect_description = db.Column(db.Text)
    reason_code_id = db.Column(db.Integer, db.ForeignKey("reason_code.id"))
    status = db.Column(db.String(80), default="Open")
    disposition = db.Column(db.String(120))
    due_date = db.Column(db.Date)
    closed_date = db.Column(db.Date)
    notes = db.Column(db.Text)
    attachment_path = db.Column(db.String(500))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    customer = db.relationship("Customer")
    department = db.relationship("Department")
    reason_code = db.relationship("ReasonCode")


class NCR(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ncr_number = db.Column(db.String(32), unique=True, nullable=False)
    date_opened = db.Column(db.Date, default=date.today)
    source = db.Column(db.String(120))
    job_order_number = db.Column(db.String(120))
    part_number = db.Column(db.String(120))
    customer_id = db.Column(db.Integer, db.ForeignKey("customer.id"))
    department_id = db.Column(db.Integer, db.ForeignKey("department.id"))
    description = db.Column(db.Text)
    requirement_violated = db.Column(db.Text)
    containment = db.Column(db.Text)
    disposition = db.Column(db.String(120))
    status = db.Column(db.String(80), default="Open")
    due_date = db.Column(db.Date)
    closed_date = db.Column(db.Date)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    customer = db.relationship("Customer")
    department = db.relationship("Department")


class DMR(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    dmr_number = db.Column(db.String(32), unique=True, nullable=False)
    date_opened = db.Column(db.Date, default=date.today)
    source_name = db.Column(db.String(120))
    po_job_order_number = db.Column(db.String(120))
    part_number = db.Column(db.String(120))
    quantity_received = db.Column(db.Integer)
    quantity_discrepant = db.Column(db.Integer)
    department_id = db.Column(db.Integer, db.ForeignKey("department.id"))
    defect_description = db.Column(db.Text)
    disposition = db.Column(db.String(120))
    material_location = db.Column(db.String(180))
    containment_action = db.Column(db.Text)
    status = db.Column(db.String(80), default="Open")
    due_date = db.Column(db.Date)
    closed_date = db.Column(db.Date)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    department = db.relationship("Department")


class Deviation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    deviation_number = db.Column(db.String(32), unique=True, nullable=False)
    request_date = db.Column(db.Date, default=date.today)
    requested_by_role = db.Column(db.String(180))
    customer_id = db.Column(db.Integer, db.ForeignKey("customer.id"))
    part_number = db.Column(db.String(120))
    order_job_number = db.Column(db.String(120))
    drawing_revision_affected = db.Column(db.String(180))
    requirement_deviated_from = db.Column(db.Text)
    reason_for_deviation = db.Column(db.Text)
    quantity_affected = db.Column(db.Integer)
    deviation_type = db.Column(db.String(80))
    effective_start_date = db.Column(db.Date)
    effective_end_date = db.Column(db.Date)
    risk_level = db.Column(db.String(80), default="Low")
    customer_approval_required = db.Column(db.Boolean, default=False)
    customer_approval_received = db.Column(db.Boolean, default=False)
    internal_approval_status = db.Column(db.String(80), default="Draft")
    approved_by_role = db.Column(db.String(180))
    approval_date = db.Column(db.Date)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    customer = db.relationship("Customer")


class AttendanceMoraleMetric(db.Model):
    """Aggregate-only morale/attendance input.

    Do not store employee names, employee numbers, payroll IDs, badge IDs, or any row that can identify a person.
    Each row should be a department/day or department/month aggregate.
    """
    id = db.Column(db.Integer, primary_key=True)
    period_date = db.Column(db.Date, default=date.today)
    department_id = db.Column(db.Integer, db.ForeignKey("department.id"))
    scheduled_headcount = db.Column(db.Integer, default=0)
    present_headcount = db.Column(db.Integer, default=0)
    overtime_hours = db.Column(db.Float, default=0)
    employees_over_50_hours = db.Column(db.Integer, default=0)
    exhausted_pto_count = db.Column(db.Integer, default=0)
    pto_absence_count = db.Column(db.Integer, default=0)
    unpaid_timeoff_count = db.Column(db.Integer, default=0)
    sick_absence_count = db.Column(db.Integer, default=0)
    turnover_count = db.Column(db.Integer, default=0)
    work_shortage_hours = db.Column(db.Float, default=0)
    staffing_shortage_count = db.Column(db.Integer, default=0)
    note = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    department = db.relationship("Department")


class ActivityLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    area = db.Column(db.String(80))
    action = db.Column(db.String(120))
    detail = db.Column(db.Text)


CSS = """
body{margin:0;background:#0b1117;color:#e8eef5;font:14px Segoe UI,Arial}
aside{position:fixed;inset:0 auto 0 0;width:215px;background:#080d12;border-right:1px solid #263647;padding:18px}
.brand{font-size:22px;font-weight:700;margin-bottom:18px;color:#fff}
nav a{display:block;color:#95a4b5;text-decoration:none;padding:9px 10px;border-radius:8px}
nav a:hover{background:#172433;color:white}
main{margin-left:215px;padding:20px;max-width:1450px}
h1{font-size:22px;margin:0 0 16px}h2{font-size:16px;color:#d7e1ec}
.cards{display:grid;grid-template-columns:repeat(5,minmax(140px,1fr));gap:12px}
.card,section{background:#121b24;border:1px solid #263647;border-radius:12px;padding:14px}
.card span{color:#95a4b5}.card b{display:block;font-size:25px;margin-top:6px;color:#fff}
table{width:100%;border-collapse:collapse;background:#121b24;border:1px solid #263647;border-radius:10px;overflow:hidden}
th,td{border-bottom:1px solid #263647;padding:8px;text-align:left;vertical-align:top}
th{color:#95a4b5;background:#0f1720;font-weight:600}
input,select,textarea{box-sizing:border-box;background:#0b1117;color:#e8eef5;border:1px solid #263647;border-radius:8px;padding:8px;width:100%}
textarea{min-height:90px}.form{display:grid;grid-template-columns:repeat(3,1fr);gap:12px}.wide{grid-column:1/-1}
.btn,button{background:#3e78a8;color:white;border:0;border-radius:8px;padding:8px 12px;text-decoration:none;display:inline-block;cursor:pointer}
.secondary{background:#4a5563}.danger{background:#8c3d3d}.muted{color:#95a4b5}.toolbar{display:flex;gap:8px;margin-bottom:12px;align-items:center;flex-wrap:wrap}
.notice{border-left:4px solid #3e78a8;background:#101923;padding:10px;border-radius:8px;margin-bottom:14px;color:#c8d4df}
.risk-low{color:#82c989}.risk-medium{color:#d6b85a}.risk-high{color:#e09054}.risk-critical{color:#f06b6b}
"""

BASE = """<!doctype html><html><head><title>ForgeQC</title><style>{{css}}</style></head>
<body><aside><div class='brand'>ForgeQC</div><nav>
<a href='/'>Dashboard</a><a href='/rma'>RMA</a><a href='/ncr'>NCR</a><a href='/dmr'>DMR</a>
<a href='/deviation'>Deviations</a><a href='/metrics'>Metrics</a><a href='/morale'>Morale Pulse</a><a href='/admin'>Admin</a>
</nav></aside><main><h1>{{title}}</h1>{{body|safe}}</main></body></html>"""


def page(title, body):
    return render_template_string(BASE, css=CSS, title=title, body=body)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or str(value).strip() == "":
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def as_int(value, default=0):
    try:
        if value in ("", None):
            return default
        return int(float(value))
    except (TypeError, ValueError):
        return default


def as_float(value, default=0.0):
    try:
        if value in ("", None):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def get_or_create(model, name):
    if name is None or str(name).strip() == "":
        return None
    clean = str(name).strip()
    row = model.query.filter(db.func.lower(model.name) == clean.lower()).first()
    if row is None:
        row = model(name=clean)
        db.session.add(row)
        db.session.flush()
    return row


def next_number(model, field_name, prefix):
    year = date.today().year
    existing = model.query.filter(getattr(model, field_name).like(f"{prefix}-{year}-%")).count()
    return f"{prefix}-{year}-{existing + 1:04d}"


def log(area, action, detail=""):
    db.session.add(ActivityLog(area=area, action=action, detail=detail))


def morale_score(metric):
    absentee_gap = max((metric.scheduled_headcount or 0) - (metric.present_headcount or 0), 0)
    score = 0
    score += min((metric.overtime_hours or 0) * 0.4, 30)
    score += (metric.employees_over_50_hours or 0) * 5
    score += (metric.exhausted_pto_count or 0) * 4
    score += (metric.pto_absence_count or 0) * 3
    score += (metric.unpaid_timeoff_count or 0) * 3
    score += (metric.sick_absence_count or 0) * 3
    score += (metric.turnover_count or 0) * 8
    score += (metric.staffing_shortage_count or 0) * 5
    score += absentee_gap * 4
    score += min((metric.work_shortage_hours or 0) * 0.25, 10)
    return round(score, 1)


def risk_label(score):
    if score >= 60:
        return "Critical"
    if score >= 35:
        return "High"
    if score >= 18:
        return "Medium"
    return "Low"


def row_count(model, closed_field="status"):
    field = getattr(model, closed_field)
    return model.query.filter(field != "Closed").count()


def ref_options(model, selected=None):
    html = "<option value=''></option>"
    for item in model.query.filter_by(active=True).order_by(model.name):
        sel = " selected" if selected and selected == item.id else ""
        html += f"<option value='{item.id}'{sel}>{item.name}</option>"
    return html


def create_app():
    app = Flask(__name__)
    root = Path(__file__).resolve().parent.parent
    imports_dir = root / "data" / "imports"
    exports_dir = root / "data" / "exports"
    imports_dir.mkdir(parents=True, exist_ok=True)
    exports_dir.mkdir(parents=True, exist_ok=True)

    app.config["SQLALCHEMY_DATABASE_URI"] = f"sqlite:///{root / 'data' / 'forgeqc.db'}"
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    db.init_app(app)

    with app.app_context():
        db.create_all()
        for model, values in [
            (Department, ["Laser", "Tube Laser", "Machining", "Welding", "Forming", "Powder", "Quality", "Shipping"]),
            (ReasonCode, ["Wrong revision", "Wrong material", "Wrong finish", "Dimensional issue", "Missing feature", "Weld defect", "Powder coat defect", "Damaged"]),
            (Customer, ["Internal"]),
        ]:
            for value in values:
                get_or_create(model, value)
        db.session.commit()

    def activity_table(limit=12):
        rows = "".join(
            f"<tr><td>{x.created_at:%m/%d/%Y %H:%M}</td><td>{x.area}</td><td>{x.action}</td><td>{x.detail or ''}</td></tr>"
            for x in ActivityLog.query.order_by(ActivityLog.id.desc()).limit(limit)
        )
        return f"<table><tr><th>Date</th><th>Area</th><th>Action</th><th>Detail</th></tr>{rows}</table>"

    def morale_summary_rows():
        rows = ""
        metrics = AttendanceMoraleMetric.query.order_by(AttendanceMoraleMetric.period_date.desc()).limit(10)
        for item in metrics:
            score = morale_score(item)
            label = risk_label(score)
            dept = item.department.name if item.department else "Unassigned"
            rows += f"<tr><td>{item.period_date}</td><td>{dept}</td><td>{score}</td><td class='risk-{label.lower()}'>{label}</td><td>{item.note or ''}</td></tr>"
        return rows or "<tr><td colspan='5' class='muted'>No aggregate morale inputs yet.</td></tr>"

    @app.route("/")
    def dashboard():
        latest_scores = [morale_score(x) for x in AttendanceMoraleMetric.query.order_by(AttendanceMoraleMetric.period_date.desc()).limit(20)]
        avg_score = round(sum(latest_scores) / len(latest_scores), 1) if latest_scores else 0
        avg_label = risk_label(avg_score)
        cards = "".join(
            f"<div class='card'><span>{k}</span><b>{v}</b></div>"
            for k, v in {
                "Open RMAs": row_count(RMA),
                "Open NCRs": row_count(NCR),
                "Open DMRs": row_count(DMR),
                "Open Deviations": Deviation.query.filter(Deviation.internal_approval_status != "Closed").count(),
                "Morale Pulse": f"{avg_label} {avg_score}",
            }.items()
        )
        body = f"""
        <div class='cards'>{cards}</div><br>
        <section><h2>Privacy-safe Morale Pulse</h2>
        <div class='notice'>Aggregate department metrics only. No employee names, numbers, badge IDs, payroll IDs, or personal attendance records are stored.</div>
        <table><tr><th>Period</th><th>Department</th><th>Score</th><th>Risk</th><th>Note</th></tr>{morale_summary_rows()}</table></section><br>
        <section><h2>Recent Activity</h2>{activity_table()}</section>
        """
        return page("Dashboard", body)

    def list_page(model, path, title, cols):
        rows = ""
        for row in model.query.order_by(model.id.desc()).limit(500):
            rows += "<tr>" + "".join(f"<td>{getattr(row, col, '') or ''}</td>" for col in cols) + f"<td><a href='/{path}/{row.id}/edit'>Edit</a></td></tr>"
        extra = ""
        if path == "rma":
            extra = "<form method='post' action='/rma/import'><button>Import Local RMA_Tracker.xlsx</button></form><a class='btn secondary' href='/rma/export'>Export</a>"
        if path == "deviation":
            extra = "<a class='btn secondary' href='/deviation/export'>Export</a>"
        heads = "".join(f"<th>{col.replace('_', ' ').title()}</th>" for col in cols)
        return page(title, f"<div class='toolbar'><a class='btn' href='/{path}/new'>New</a>{extra}</div><table><tr>{heads}<th></th></tr>{rows}</table>")

    app.add_url_rule("/rma", "rma_list", lambda: list_page(RMA, "rma", "RMA", ["rma_number", "date_opened", "order_number", "part_number", "responsible_area", "status"]))
    app.add_url_rule("/ncr", "ncr_list", lambda: list_page(NCR, "ncr", "NCR", ["ncr_number", "date_opened", "source", "part_number", "status"]))
    app.add_url_rule("/dmr", "dmr_list", lambda: list_page(DMR, "dmr", "DMR", ["dmr_number", "date_opened", "source_name", "part_number", "status"]))
    app.add_url_rule("/deviation", "dev_list", lambda: list_page(Deviation, "deviation", "Deviation Requests", ["deviation_number", "request_date", "requested_by_role", "part_number", "risk_level", "internal_approval_status"]))

    def record_form(kind, row):
        title = getattr(row, f"{kind}_number", getattr(row, "deviation_number", "New"))
        customer_select = ref_options(Customer, getattr(row, "customer_id", None))
        department_select = ref_options(Department, getattr(row, "department_id", None))
        reason_select = ref_options(ReasonCode, getattr(row, "reason_code_id", None))
        description_value = getattr(row, "defect_description", None) or getattr(row, "description", None) or getattr(row, "requirement_deviated_from", None) or ""
        person_value = getattr(row, "responsible_area", None) or getattr(row, "requested_by_role", None) or ""
        order_value = getattr(row, "order_number", None) or getattr(row, "job_order_number", None) or getattr(row, "po_job_order_number", None) or getattr(row, "order_job_number", None) or getattr(row, "source", None) or getattr(row, "source_name", None) or ""
        status_value = getattr(row, "status", None) or getattr(row, "internal_approval_status", None) or "Open"
        body = f"""
        <form method='post' class='form'>
        <h2 class='wide'>{title}</h2>
        <label>Date<input type='date' name='date' value='{getattr(row, "date_opened", None) or getattr(row, "request_date", None) or ""}'></label>
        <label>Customer<select name='customer_id'>{customer_select}</select></label>
        <label>Department<select name='department_id'>{department_select}</select></label>
        <label>Reason Code<select name='reason_code_id'>{reason_select}</select></label>
        <label>Order / Job / Source<input name='order' value='{order_value or ""}'></label>
        <label>Part Number<input name='part_number' value='{getattr(row, "part_number", "") or ""}'></label>
        <label>Responsible Area / Role<input name='person' value='{person_value or ""}'></label>
        <label>Status<input name='status' value='{status_value or ""}'></label>
        <label>Risk Level<select name='risk_level'><option>Low</option><option>Medium</option><option>High</option><option>Critical</option></select></label>
        <label class='wide'>Description / Requirement<textarea name='description'>{description_value}</textarea></label>
        <label class='wide'>Notes<textarea name='notes'>{getattr(row, "notes", "") or ""}</textarea></label>
        <button>Save</button>
        </form>
        """
        return page(f"{kind.upper()} Form", body)

    def save_record(kind, row):
        form = request.form
        if hasattr(row, "date_opened"):
            row.date_opened = parse_date(form.get("date")) or row.date_opened
        if hasattr(row, "request_date"):
            row.request_date = parse_date(form.get("date")) or row.request_date
        if hasattr(row, "customer_id"):
            row.customer_id = as_int(form.get("customer_id"), None)
        if hasattr(row, "department_id"):
            row.department_id = as_int(form.get("department_id"), None)
        if hasattr(row, "reason_code_id"):
            row.reason_code_id = as_int(form.get("reason_code_id"), None)
        if hasattr(row, "part_number"):
            row.part_number = form.get("part_number")
        if hasattr(row, "status"):
            row.status = form.get("status") or "Open"
        if hasattr(row, "internal_approval_status"):
            row.internal_approval_status = form.get("status") or "Draft"
        if hasattr(row, "risk_level"):
            row.risk_level = form.get("risk_level") or row.risk_level
        if kind == "rma":
            row.order_number = form.get("order")
            row.responsible_area = form.get("person")
            row.defect_description = form.get("description")
        elif kind == "ncr":
            row.source = form.get("order")
            row.description = form.get("description")
        elif kind == "dmr":
            row.source_name = form.get("order")
            row.defect_description = form.get("description")
        elif kind == "deviation":
            row.order_job_number = form.get("order")
            row.requested_by_role = form.get("person")
            row.requirement_deviated_from = form.get("description")
        row.notes = form.get("notes")

    def edit_record(model, kind, field, prefix, path, row_id=None):
        row = model.query.get(row_id) if row_id else model(**{field: next_number(model, field, prefix)})
        if row is None:
            abort(404)
        if request.method == "POST":
            save_record(kind, row)
            db.session.add(row)
            log(kind.upper(), "Save", getattr(row, field))
            db.session.commit()
            return redirect("/" + path)
        return record_form(kind, row)

    for url, model, kind, field, prefix in [
        ("/rma", RMA, "rma", "rma_number", "RMA"),
        ("/ncr", NCR, "ncr", "ncr_number", "NCR"),
        ("/dmr", DMR, "dmr", "dmr_number", "DMR"),
        ("/deviation", Deviation, "deviation", "deviation_number", "DEV"),
    ]:
        app.add_url_rule(url + "/new", kind + "_new", lambda model=model, kind=kind, field=field, prefix=prefix, url=url: edit_record(model, kind, field, prefix, url.strip("/")), methods=["GET", "POST"])
        app.add_url_rule(url + "/<int:row_id>/edit", kind + "_edit", lambda row_id, model=model, kind=kind, field=field, prefix=prefix, url=url: edit_record(model, kind, field, prefix, url.strip("/"), row_id), methods=["GET", "POST"])

    @app.route("/rma/import", methods=["POST"])
    def import_rma():
        path = imports_dir / "RMA_Tracker.xlsx"
        if not path.exists():
            log("RMA", "Import failed", "Local data/imports/RMA_Tracker.xlsx not found")
            db.session.commit()
            return redirect("/rma")

        workbook = load_workbook(path, data_only=True)
        imported = 0
        skipped = 0

        for sheet_name, model in [("Customer List", Customer), ("Department List", Department), ("Reason Codes", ReasonCode)]:
            if sheet_name in workbook.sheetnames:
                for (value,) in workbook[sheet_name].iter_rows(min_row=2, max_col=1, values_only=True):
                    get_or_create(model, value)

        if "RMA Log" not in workbook.sheetnames:
            log("RMA", "Import failed", "RMA Log sheet missing")
            db.session.commit()
            return redirect("/rma")

        for source_row in workbook["RMA Log"].iter_rows(min_row=2, values_only=True):
            opened, order, customer, department, person, description, reason = (list(source_row) + [None] * 7)[:7]
            if not any([opened, order, customer, department, person, description, reason]):
                skipped += 1
                continue

            opened_date = parse_date(opened)
            cust = get_or_create(Customer, customer)
            dept = get_or_create(Department, department)
            reason_code = get_or_create(ReasonCode, reason)

            duplicate = RMA.query.filter_by(
                date_opened=opened_date,
                order_number=str(order).strip() if order else None,
                department_id=dept.id if dept else None,
                defect_description=str(description).strip() if description else None,
            ).first()
            if duplicate:
                skipped += 1
                continue

            db.session.add(RMA(
                rma_number=next_number(RMA, "rma_number", "RMA"),
                date_opened=opened_date,
                customer_id=cust.id if cust else None,
                order_number=str(order).strip() if order else None,
                department_id=dept.id if dept else None,
                responsible_area=str(person).strip() if person else None,
                defect_description=str(description).strip() if description else None,
                reason_code_id=reason_code.id if reason_code else None,
                status="Open",
            ))
            imported += 1

        log("RMA", "Import", f"Imported {imported}; skipped {skipped}. Workbook stayed local.")
        db.session.commit()
        return redirect("/rma")

    def export_xlsx(filename, headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for item in rows:
            sheet.append(item)
        output = exports_dir / filename
        workbook.save(output)
        return send_file(output, as_attachment=True)

    @app.route("/rma/export")
    def export_rma():
        return export_xlsx("rma_export.xlsx", ["RMA", "Date", "Order", "Part", "Responsible Area", "Status"], [
            [r.rma_number, r.date_opened, r.order_number, r.part_number, r.responsible_area, r.status] for r in RMA.query.all()
        ])

    @app.route("/deviation/export")
    def export_dev():
        return export_xlsx("deviation_export.xlsx", ["Deviation", "Date", "Part", "Risk", "Status"], [
            [r.deviation_number, r.request_date, r.part_number, r.risk_level, r.internal_approval_status] for r in Deviation.query.all()
        ])

    @app.route("/admin", methods=["GET", "POST"])
    def admin():
        models = {"customers": Customer, "departments": Department, "reason_codes": ReasonCode}
        if request.method == "POST":
            model = models[request.form["table"]]
            name = request.form.get("name", "").strip()
            if name:
                row = get_or_create(model, name)
                row.description = request.form.get("description")
                row.active = True
                log("Admin", "Reference save", f"{request.form['table']}: {name}")
                db.session.commit()
            return redirect("/admin")

        body = ""
        for name, model in models.items():
            rows = "".join(f"<tr><td>{r.name}</td><td>{r.description or ''}</td><td>{'Yes' if r.active else 'No'}</td></tr>" for r in model.query.order_by(model.name))
            body += f"""
            <section><h2>{name.replace('_', ' ').title()}</h2>
            <form method='post' class='form'><input type='hidden' name='table' value='{name}'>
            <label>Name<input name='name'></label><label>Description<input name='description'></label><div><br><button>Add / Reactivate</button></div></form>
            <table><tr><th>Name</th><th>Description</th><th>Active</th></tr>{rows}</table></section><br>
            """
        return page("Admin Reference Data", body)

    @app.route("/morale", methods=["GET", "POST"])
    def morale():
        if request.method == "POST":
            metric = AttendanceMoraleMetric(
                period_date=parse_date(request.form.get("period_date")) or date.today(),
                department_id=as_int(request.form.get("department_id"), None),
                scheduled_headcount=as_int(request.form.get("scheduled_headcount")),
                present_headcount=as_int(request.form.get("present_headcount")),
                overtime_hours=as_float(request.form.get("overtime_hours")),
                employees_over_50_hours=as_int(request.form.get("employees_over_50_hours")),
                exhausted_pto_count=as_int(request.form.get("exhausted_pto_count")),
                pto_absence_count=as_int(request.form.get("pto_absence_count")),
                unpaid_timeoff_count=as_int(request.form.get("unpaid_timeoff_count")),
                sick_absence_count=as_int(request.form.get("sick_absence_count")),
                turnover_count=as_int(request.form.get("turnover_count")),
                work_shortage_hours=as_float(request.form.get("work_shortage_hours")),
                staffing_shortage_count=as_int(request.form.get("staffing_shortage_count")),
                note=request.form.get("note"),
            )
            db.session.add(metric)
            log("Morale", "Aggregate pulse input", "Privacy-safe department metric saved")
            db.session.commit()
            return redirect("/morale")

        dept_select = ref_options(Department)
        rows = ""
        for metric in AttendanceMoraleMetric.query.order_by(AttendanceMoraleMetric.period_date.desc(), AttendanceMoraleMetric.id.desc()).limit(100):
            score = morale_score(metric)
            label = risk_label(score)
            dept = metric.department.name if metric.department else "Unassigned"
            rows += f"<tr><td>{metric.period_date}</td><td>{dept}</td><td>{metric.scheduled_headcount}</td><td>{metric.present_headcount}</td><td>{metric.overtime_hours}</td><td>{metric.sick_absence_count}</td><td>{metric.turnover_count}</td><td>{score}</td><td class='risk-{label.lower()}'>{label}</td></tr>"

        body = f"""
        <div class='notice'>Privacy rule: enter aggregate department metrics only. No employee names, employee numbers, badge IDs, payroll IDs, schedules, or individual attendance records.</div>
        <form method='post' class='form'>
        <label>Period Date<input type='date' name='period_date'></label>
        <label>Department<select name='department_id'>{dept_select}</select></label>
        <label>Scheduled Headcount<input name='scheduled_headcount' type='number' min='0'></label>
        <label>Present Headcount<input name='present_headcount' type='number' min='0'></label>
        <label>Overtime Hours<input name='overtime_hours' type='number' step='0.1' min='0'></label>
        <label>Employees Over 50 Hours<input name='employees_over_50_hours' type='number' min='0'></label>
        <label>Exhausted PTO Count<input name='exhausted_pto_count' type='number' min='0'></label>
        <label>PTO Absence Count<input name='pto_absence_count' type='number' min='0'></label>
        <label>Unpaid Timeoff Count<input name='unpaid_timeoff_count' type='number' min='0'></label>
        <label>Sick Absence Count<input name='sick_absence_count' type='number' min='0'></label>
        <label>Turnover Count<input name='turnover_count' type='number' min='0'></label>
        <label>Work Shortage Hours<input name='work_shortage_hours' type='number' step='0.1' min='0'></label>
        <label>Staffing Shortage Count<input name='staffing_shortage_count' type='number' min='0'></label>
        <label class='wide'>Operational Note<textarea name='note'></textarea></label>
        <button>Save Aggregate Pulse Input</button>
        </form><br>
        <section><h2>Morale Pulse History</h2><table><tr><th>Period</th><th>Department</th><th>Scheduled</th><th>Present</th><th>OT Hours</th><th>Sick</th><th>Turnover</th><th>Score</th><th>Risk</th></tr>{rows}</table></section>
        """
        return page("Morale Pulse Compiler", body)

    @app.route("/metrics")
    def metrics():
        customer_rows = "".join(
            f"<tr><td>{name or 'Unassigned'}</td><td>{count}</td></tr>"
            for name, count in db.session.query(Customer.name, db.func.count(RMA.id)).outerjoin(RMA).group_by(Customer.id).order_by(db.func.count(RMA.id).desc()).all()
        )

        dept_quality = defaultdict(lambda: {"rma": 0, "ncr": 0, "dmr": 0, "pulse": []})
        for dept, count in db.session.query(Department.name, db.func.count(RMA.id)).outerjoin(RMA).group_by(Department.id):
            dept_quality[dept or "Unassigned"]["rma"] = count
        for dept, count in db.session.query(Department.name, db.func.count(NCR.id)).outerjoin(NCR).group_by(Department.id):
            dept_quality[dept or "Unassigned"]["ncr"] = count
        for dept, count in db.session.query(Department.name, db.func.count(DMR.id)).outerjoin(DMR).group_by(Department.id):
            dept_quality[dept or "Unassigned"]["dmr"] = count
        for metric in AttendanceMoraleMetric.query.all():
            dept = metric.department.name if metric.department else "Unassigned"
            dept_quality[dept]["pulse"].append(morale_score(metric))

        corr_rows = ""
        for dept, values in sorted(dept_quality.items()):
            avg = round(sum(values["pulse"]) / len(values["pulse"]), 1) if values["pulse"] else 0
            corr_rows += f"<tr><td>{dept}</td><td>{values['rma']}</td><td>{values['ncr']}</td><td>{values['dmr']}</td><td>{avg}</td><td>{risk_label(avg)}</td></tr>"

        body = f"""
        <section><h2>RMA Pareto by Customer</h2><table><tr><th>Customer</th><th>RMA Count</th></tr>{customer_rows}</table></section><br>
        <section><h2>Quality vs Aggregate Morale Pulse</h2>
        <div class='notice'>Correlation is department-level only. It does not store or display individual employee data.</div>
        <table><tr><th>Department</th><th>RMA</th><th>NCR</th><th>DMR</th><th>Average Pulse Score</th><th>Risk</th></tr>{corr_rows}</table></section>
        """
        return page("Metrics", body)

    return app
